import openpyxl
from pathlib import Path

FILE_PATH = Path(r"d:\Project PKL\qa-reports\whitebox\LAPORAN-WHITEBOX.xlsx")

def validate():
    print(f"Loading {FILE_PATH}...")
    wb = openpyxl.load_workbook(FILE_PATH, data_only=False)
    
    expected_sheets = [
        "Ringkasan",
        "Hasil Test",
        "Coverage",
        "Temuan",
        "Dependency & Secret",
        "Kompleksitas & Duplikasi",
        "Rekomendasi Test",
        "Keterbatasan"
    ]
    
    print(f"Sheet names: {wb.sheetnames}")
    assert wb.sheetnames == expected_sheets, f"Expected {expected_sheets}, got {wb.sheetnames}"
    
    error_tokens = ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!"]
    errors_found = []
    
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"Checking sheet '{name}': {ws.max_row} rows, {ws.max_column} cols, freeze={ws.freeze_panes}")
        
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                val = str(cell.value) if cell.value is not None else ""
                for err in error_tokens:
                    if err in val:
                        errors_found.append((name, cell.coordinate, val))
                        
    if errors_found:
        print(f"ERROR: Found {len(errors_found)} error cells: {errors_found}")
        return False
        
    print("No formula syntax errors found across all sheets!")
    
    # Inspect Ringkasan formulas
    ws1 = wb["Ringkasan"]
    print("Ringkasan formulas:")
    print("  B5 (Total Test):", ws1["B5"].value)
    print("  B6 (Pass):", ws1["B6"].value)
    print("  B7 (Fail):", ws1["B7"].value)
    print("  B8 (Skip):", ws1["B8"].value)
    print("  B9 (Rate):", ws1["B9"].value)
    print("  F5 (Critical):", ws1["F5"].value)
    print("  F6 (High):", ws1["F6"].value)
    print("  F7 (Medium):", ws1["F7"].value)
    print("  F8 (Low):", ws1["F8"].value)
    print("  F9 (Total Temuan):", ws1["F9"].value)

    print("ALL VALIDATION CHECKS PASSED SUCCESSFULLY!")
    return True

if __name__ == '__main__':
    validate()
